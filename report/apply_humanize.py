import copy, re, sys
sys.path.insert(0,'.')
from humanize_rewrites import R
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
d=Document("PharmaCompare_Project_Report_before_humanize.docx")
ps=d.paragraphs
def apply(p,text):
    runs=[r for r in p.runs if r.text]
    plain=next((r for r in runs if not r.bold),None)
    tmpl_r=copy.deepcopy((plain or runs[0])._r)
    for t in tmpl_r.findall(qn('w:t')): tmpl_r.remove(t)
    rpr=tmpl_r.find(qn('w:rPr'))
    if rpr is not None:
        for tag in ('w:b','w:bCs'):
            for e in rpr.findall(qn(tag)): rpr.remove(e)
    for r in p._p.findall(qn('w:r')): p._p.remove(r)
    for part in re.split(r"(\*\*.+?\*\*)",text):
        if not part: continue
        bold=part.startswith("**") and part.endswith("**")
        t=part[2:-2] if bold else part
        r=copy.deepcopy(tmpl_r)
        rp=r.find(qn('w:rPr'))
        if bold:
            if rp is None:
                rp=OxmlElement('w:rPr'); r.insert(0,rp)
            rp.append(OxmlElement('w:b'))
        te=OxmlElement('w:t'); te.text=t; te.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); r.append(te)
        p._p.append(r)
for i,txt in R.items(): apply(ps[i],txt)
for row in d.tables[36].rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs:
                if "12,007" in r.text: r.text=r.text.replace("12,007","11,958")
d.save("PharmaCompare_Project_Report.docx")
print("rewrote",len(R))
