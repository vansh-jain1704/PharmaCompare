import re, sys, pypdfium2 as pdfium
from docx import Document
d=Document("PharmaCompare_Project_Report.docx")
def pages(path):
    pdf=pdfium.PdfDocument(path)
    return [pdf[i].get_textpage().get_text_range().replace("\r","") for i in range(len(pdf))]
def printed(t):
    lines=[l.strip() for l in t.split("\n") if l.strip()]
    for l in reversed(lines[-4:]):
        if re.fullmatch(r"[ivxlc]+|\d{1,3}",l,re.I): return l.lower()
def find(pgs,pat,start):
    for i in range(start,len(pgs)):
        if re.search(pat,pgs[i]): return i
entries=[]
for ti in (0,1):
    for r in d.tables[ti].rows[1:]:
        c=[x.text.strip() for x in r.cells]
        if c[1] and c[2]: entries.append(("sec",c[0],c[1],c[2]))
for ti,kind in ((2,"fig"),(3,"tab")):
    for r in d.tables[ti].rows[1:]:
        c=[x.text.strip() for x in r.cells]
        if c[0] and c[2]: entries.append((kind,c[0],c[1],c[2]))
for label,path in (("BEFORE","C:/Users/vikra/AppData/Local/Temp/rb/before.pdf"),("AFTER","D:/PharmaCompare/report/PharmaCompare_Project_Report.pdf")):
    pgs=pages(path); start=find(pgs,r"CHAPTER 1\s*\n\s*PREAMBLE",0) or 0
    bad=[]
    for kind,no,title,pg in entries:
        if kind=="sec":
            if not re.match(r"\d",title) and not no: continue
            pat=r"CHAPTER\s*%s\s*\n"%no if no.isdigit() else re.escape(title.split(" ",1)[1])
        elif kind=="fig": pat=r"Figure\s*%s[\s:.]"%re.escape(no)
        else: pat=r"Table\s*%s[\s:.]"%re.escape(no)
        i=find(pgs,pat,start); pr=printed(pgs[i]) if i is not None else None
        if pr!=pg.lower(): bad.append((kind,no,title[:28],pg,pr))
    print(label,len(pgs),"pages; mismatches:",len(bad),bad[:10])
